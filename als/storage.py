"""The data file: an Excel workbook that doctors can edit in Excel as well as in the app.

Layout (sheet: columns, see als.model.COLUMNS); every value is text such as "03/2026":
  README           how to edit the file (written by the app, not read)
  forms            one row per patient
  treatments       one row per treatment course
  treatment_types  the treatments the app offers
Downloads add result sheets (summary, score_timeline); they are ignored when a workbook is read.

read_workbook() refuses a workbook whose sheets or columns differ from this layout (WorkbookError), because
saving it again would silently drop data. save() writes a complete new file and then swaps it in, so the file on
disk is never half written.
"""
from __future__ import annotations

import contextlib
import io
import os
import tempfile
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .model import BOX_COLS, COLUMNS, FORMS, TREATMENTS, TYPES, Tables, empty_tables, normalize

README = "README"
RESULT_SHEETS = ("summary", "score_timeline")       # in downloads only; never read back
EDIT_ROWS = 2000                                    # rows prepared for typing in Excel (text format, dropdowns)
TEXT = "@"                                          # Excel number format "Text": 03/2026 stays text

README_LINES = [
    "ALS progress data (ALSFRS-R). Edit it in the app or here in Excel: both use this file.",
    "",
    "Editing in Excel",
    "  - Do not rename, add or delete sheets or columns: the app then refuses to use the file until it is fixed.",
    "  - While this file is open in Excel the app cannot save. Close it in Excel when you are done;",
    "    the app shows your changes the next time you click in it.",
    "  - Formatting and colours you add are not kept: the app rewrites the file in this layout when it saves.",
    "  - Mistakes (wrong month, unknown drug, ...) are listed in the app: All patients & data, column check.",
    "",
    "forms: one row per patient. patient_id = study code only (MND-13), never a name or HN.",
    "  q{n}_{item}_L3 / L2 / L1 / L0 = month (MM/YYYY, AD) the patient first reached that level.",
    "  Leave level 4 (normal) out. YYYY alone = month unknown (counted as July). ? = reached by the latest month (the patient has reached this level, but we don’t know the month).",
    "  last_assessed = month of the latest assessment (MM/YYYY), if later than the last change.",
    "treatments: one row per treatment course: patient_id, treatment, start, stop (MM/YYYY; blank stop = still on).",
    "  A patient can have any number of courses, also of the same drug (stopped and restarted).",
    "treatment_types: the treatments on offer. Add a row to add a new drug.",
    "  Rename a drug in the app (Treatments page): it renames the drug in every course too.",
    "",
    "Cells are formatted as Text so that 03/2026 stays 03/2026. If Excel still turns it into a date, the app reads",
    "it back as that month. Scores, slopes and the monthly timeline are calculated by the app.",
]
MONTH_PROMPT = "MM/YYYY in AD, e.g. 03/2026"
BOX_PROMPT = ("Month the patient first reached this level: MM/YYYY, YYYY if only the year is known, "
              "or ? (reached by the latest month (the patient has reached this level, but we don’t know the month)). Blank = not reached.")
HEADER_FONT = Font(name="Arial", size=10, bold=True)
HEADER_FILL = PatternFill("solid", fgColor="FFD9E2F3")
# Rows and columns kept in view while scrolling (Excel "freeze panes"); set once per sheet, because openpyxl
# adds a pane selection on every assignment and Excel rejects the extra ones ("We found a problem with some content").
FREEZE = {TYPES: "A2"}                              # header row only; other sheets also keep column A in view
WIDTHS = {"patient_id": 12, "form_language": 9, "last_assessed": 12, "notes": 40, "treatment": 20,
          "start": 10, "stop": 10}


class WorkbookError(ValueError):
    """The workbook cannot be used as it is; `problems` says what to fix."""

    def __init__(self, problems: list[str]):
        super().__init__("; ".join(problems))
        self.problems = problems


# ----------------------------------------------------------------------------- reading
def read_workbook(source) -> Tables:
    """The three tables from an .xlsx path or uploaded file.

    Raises WorkbookError if the file is not a readable workbook, or if a data sheet or column is missing,
    unknown or doubled.
    """
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as err:
        raise WorkbookError([f"not a readable .xlsx workbook ({err})"]) from err
    try:
        problems = [f"unknown sheet '{name}'" for name in wb.sheetnames
                    if name not in COLUMNS and name != README and name not in RESULT_SHEETS]
        tables = empty_tables()
        for name in COLUMNS:
            if name not in wb.sheetnames:
                problems.append(f"sheet '{name}' is missing")
                continue
            table, sheet_problems = _read_sheet(wb[name], name)
            problems += sheet_problems
            tables[name] = table
    finally:
        wb.close()
    if problems:
        raise WorkbookError(problems)
    return tables


def _read_sheet(ws: Worksheet, name: str) -> tuple[pd.DataFrame, list[str]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame(columns=COLUMNS[name]), [f"sheet '{name}' has no header row"]
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    named = [h for h in header if h]
    problems = [f"{name}: column '{col}' is missing" for col in COLUMNS[name] if col not in named]
    problems += [f"{name}: unknown column '{col}'" for col in dict.fromkeys(named) if col not in COLUMNS[name]]
    problems += [f"{name}: column '{col}' appears twice" for col in dict.fromkeys(named) if named.count(col) > 1]
    keep = [i for i, h in enumerate(header) if h in COLUMNS[name]]
    data = [[_month_text(row[i]) if i < len(row) else None for i in keep] for row in rows[1:]]
    return normalize(name, pd.DataFrame(data, columns=[header[i] for i in keep])), problems


def _month_text(value):
    """A cell that Excel turned into a date (03/2026 -> 1 Mar 2026) is read back as 'MM/YYYY'."""
    return value.strftime("%m/%Y") if hasattr(value, "strftime") else value


# ----------------------------------------------------------------------------- writing
def to_workbook_bytes(tables: Tables, results: dict[str, pd.DataFrame] | None = None) -> bytes:
    """The workbook: README, the three data tables (text cells, ready for editing in Excel), then optional
    result sheets (e.g. summary, score_timeline)."""
    wb = Workbook()
    _write_readme(wb.active)
    for name in COLUMNS:
        ws = wb.create_sheet(name)
        _write_table(ws, normalize(name, tables.get(name)), as_text=True, freeze=FREEZE.get(name, "B2"))
        _prepare_for_excel(ws, name)
    for name, df in (results or {}).items():
        _write_table(wb.create_sheet(name), df, as_text=False)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _write_readme(ws: Worksheet) -> None:
    ws.title = README
    for i, line in enumerate(README_LINES, 1):
        ws.cell(i, 1, line).font = Font(name="Arial", size=10, bold=i == 1 or line == "Editing in Excel")
    ws.column_dimensions["A"].width = 120


def _write_table(ws: Worksheet, df: pd.DataFrame, as_text: bool, freeze: str = "B2") -> None:
    for j, col in enumerate(df.columns, 1):
        cell = ws.cell(1, j, col)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = WIDTHS.get(col, max(10, min(40, len(str(col)) + 2)))
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, value in enumerate(row, 1):
            if value is None or (not isinstance(value, str) and pd.isna(value)):
                continue
            cell = ws.cell(i, j, str(value) if as_text else value)
            if as_text:
                cell.number_format = TEXT
            elif isinstance(value, pd.Timestamp):
                cell.number_format = "mm/yyyy"
    ws.freeze_panes = freeze
    ws.row_dimensions[1].height = 30


def _prepare_for_excel(ws: Worksheet, name: str) -> None:
    """Text format for whole columns, dropdowns for languages / patients / drugs, and input hints."""
    cols = COLUMNS[name]

    def cells(col: str) -> str:
        letter = get_column_letter(cols.index(col) + 1)
        return f"{letter}2:{letter}{EDIT_ROWS}"

    for j in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].number_format = TEXT
    if name == FORMS:
        _dropdown(ws, '"TH,EN"', cells("form_language"), "TH or EN")
        _hint(ws, MONTH_PROMPT, cells("last_assessed"))
        first, last = cells(BOX_COLS[0]).split(":")[0], cells(BOX_COLS[-1]).split(":")[1]
        _hint(ws, BOX_PROMPT, f"{first}:{last}")
    elif name == TREATMENTS:
        _dropdown(ws, f"{FORMS}!$A$2:$A${EDIT_ROWS}", cells("patient_id"), "Not a patient_id on the forms sheet")
        _dropdown(ws, f"{TYPES}!$A$2:$A${EDIT_ROWS}", cells("treatment"),
                  "Not on the treatment_types sheet: add the drug there first")
        _hint(ws, MONTH_PROMPT, cells("start"))
        _hint(ws, MONTH_PROMPT + ". Blank = still on.", cells("stop"))


def _dropdown(ws: Worksheet, source: str, cells: str, error: str) -> None:
    """A dropdown list; other values are allowed after a warning (the app's checks report them)."""
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                        errorStyle="warning", error=error)
    dv.add(cells)
    ws.add_data_validation(dv)


def _hint(ws: Worksheet, text: str, cells: str) -> None:
    """A note Excel shows when one of these cells is selected (any value allowed)."""
    dv = DataValidation(allow_blank=True, showInputMessage=True, prompt=text)
    dv.add(cells)
    ws.add_data_validation(dv)


# ----------------------------------------------------------------------------- the data file
def load(path: Path) -> Tables:
    """The tables in the data file (empty tables if it does not exist yet). Raises WorkbookError."""
    return read_workbook(path) if Path(path).exists() else empty_tables()


def save(path: Path, tables: Tables) -> None:
    """Rewrite the data file, all or nothing.

    The new workbook goes to a temporary file next to it, which then replaces the old file in one step, so a crash
    or power cut never leaves a half-written file. If anything fails, the old file is unchanged. On Windows a
    PermissionError means the file is open in another program (Excel locks the files it has open).
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = to_workbook_bytes(tables)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix="tmp", suffix=".xlsx")
    try:
        with os.fdopen(fd, "wb") as fh:
            fh.write(data)
            fh.flush()
            os.fsync(fh.fileno())
        os.replace(tmp, path)
    except BaseException:
        with contextlib.suppress(OSError):
            os.remove(tmp)
        raise


def file_version(path: Path) -> tuple[int, int] | None:
    """(modification time, size) of the file, None if it does not exist. Changes whenever the file is saved."""
    try:
        stat = Path(path).stat()
    except FileNotFoundError:
        return None
    return stat.st_mtime_ns, stat.st_size


def open_in_excel(path: Path) -> bool:
    """True while Excel has the file open (Excel keeps a hidden '~$<name>' file next to it)."""
    path = Path(path)
    return (path.parent / f"~${path.name}").exists()
